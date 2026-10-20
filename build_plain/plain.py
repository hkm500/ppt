#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""plain.py — 按《数字货币3.pptx》的朴素风格生成幻灯片。

以 数字货币3.pptx 为母版（保留封面，删除其余页），用 python-pptx + 原生 XML
复刻其版式：白底、左对齐楷体粗体标题（40pt，缩放 90%）、Times New Roman / 楷体
正文（20–28pt）、C00000 红色粗体强调、1F3864 表头的简洁表格、图片 + 说明文字。
无卡片、图标、色块、章节分隔页。

公共接口（见 content/*.py 的用法）：

    d.text(title, items, notes=None, ...)              文字页
    d.table(title, header, rows, widths=None, foot=None, notes=None, ...)   表格页
    d.pic(title, img, items=None, notes=None, layout='top'|'left'|'right')   图片页
    d.cover(title, notes)                               封面（重设标题）

items 中每个元素是一行（段）；行首标记决定样式：
    "正文"        普通段落（无项目符号）
    "# 小标题"    粗体段落
    "- 要点"      一级圆点项目符号
    "#- 要点"     粗体一级圆点（后面常跟二级要点）
    "-- 子要点"   二级短横项目符号（字号小一档）
    "!! 文字"     整段红色粗体（“思考/启示/主线/关键问题”等收束句）
    ""            空行（间隔）
行内标记：
    [[文字]]      红色粗体强调
    **文字**      粗体（黑）
    ^{上标} _{下标}   上标/下标
    `代码`        等宽字体（哈希、地址等）
"""
import copy
import datetime
import math
import os
import re
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = os.path.join(ROOT, "数字货币3.pptx")
ASSETS = os.path.join(ROOT, "build", "assets")

NSA = "http://schemas.openxmlformats.org/drawingml/2006/main"
NSP = "http://schemas.openxmlformats.org/presentationml/2006/main"
NSR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSDECL = f'xmlns:a="{NSA}" xmlns:p="{NSP}" xmlns:r="{NSR}"'

LATIN = "Times New Roman"
MONO = "Courier New"
EA = "楷体"
RED = "C00000"
NAVY = "1F3864"
GRID = "BFBFBF"
BAND = "F2F2F2"

SLIDE_W, SLIDE_H = 9144000, 6858000
BODY_X, BODY_Y, BODY_W, BODY_H = 152516, 1066862, 8915166, 5562454
INS_LR, INS_TB = 91440, 45720  # 文本框内边距（EMU）
EMU_PT = 12700.0
LINE_K = 1.2  # 行高 = 字号 × LINE_K

SUB_SIZE = {30: 24, 28: 22, 26: 20, 24: 20, 22: 18, 20: 18, 18: 16, 16: 15}


# --------------------------------------------------------------------------
# 文本度量（估算折行与高度；真实渲染以 LibreOffice / PowerPoint 为准）
# --------------------------------------------------------------------------
_NARROW = set("iIjl")
_THIN = set("frt")
_WIDE = set("mwMW")
_PUNCT_SMALL = set(".,:;'!|")
_PUNCT_MID = set("()[]-/")
_SYM_MID = set("+=<>&@#~^*_")
_HALF_SYMS = set("×÷±≈≠≤≥·°′″²³¹")
_QUOTES = set("“”‘’")
_NO_BREAK_BEFORE = set("，。、；：！？）》」』”’】〕…—·,.;:!?)]}%")
_NO_BREAK_AFTER = set("（《「『“‘【〔([{")

# 两种度量模式：
#   "ppt"：PowerPoint/WPS 近似——中西文之间不加间距，弯引号按全角 1em；
#   "lo" ：LibreOffice 近似——中西文交界自动加约 0.25em 间距，弯引号用西文字形 0.5em。
# 排版时两种模式都要“不溢出、无孤字”，这样换软件打开也稳。
MODES = ("ppt", "lo")
AUTOSPACE = 0.25


_MONO_ALL = False   # 段内含 `代码` 时，估算时把西文统一按等宽 0.6em 算（保守）


def char_em(ch, mode="ppt"):
    """字符宽度（以 em 计）。CJK 与全角标点按 1em，西文按 Times New Roman 近似。"""
    o = ord(ch)
    if ch == " ":
        return 0.25
    if _MONO_ALL and o < 0x80:
        return 0.6
    if o < 0x80:
        if ch.isdigit():
            return 0.5
        if ch == "%":
            return 0.83
        if ch in _NARROW:
            return 0.28
        if ch in _THIN:
            return 0.36
        if ch in _WIDE:
            return 0.8
        if ch.isupper():
            return 0.68
        if ch.islower():
            return 0.47
        if ch in _PUNCT_SMALL:
            return 0.27
        if ch in _PUNCT_MID:
            return 0.34
        if ch in _SYM_MID:
            return 0.6
        return 0.45
    if ch in _QUOTES:
        return 0.5 if mode == "lo" else 1.0
    if ch in _HALF_SYMS:
        return 0.55
    return 1.0


def _is_latin(ch):
    return (ch.isascii() and ch.isalnum()) or ch in "%$"


def _tokens(text, bold=False, mode="ppt"):
    """把文本切成不可再分的折行单元：[(字符串, 宽度em)]。粗体仅西文字符变宽。"""
    toks = []
    buf = ""
    bw = 0.0

    def flush():
        nonlocal buf, bw
        if buf:
            toks.append((buf, bw))
            buf, bw = "", 0.0

    for i, ch in enumerate(text):
        o = ord(ch)
        w = char_em(ch, mode) * (1.07 if (bold and o < 0x80) else 1.0)
        if o < 0x80 and ch != " " and ch not in _NO_BREAK_AFTER:
            buf += ch
            bw += w
            prv = text[i - 1] if i > 0 else ""
            nxt = text[i + 1] if i + 1 < len(text) else ""
            if nxt.isascii() and nxt.isalpha() and (
                    (ch == "-" and prv.isascii() and prv.isalpha()) or (ch == "/" and mode == "lo")):
                flush()  # 连字符 / 斜杠之后允许断行
            continue
        if ch == " ":
            flush()
            toks.append((" ", w))
            continue
        # CJK 或其他非 ASCII 字符：独立单元，但禁则处理
        flush()
        if ch in _NO_BREAK_BEFORE and toks and toks[-1][0] != " ":
            prev, pw = toks.pop()
            toks.append((prev + ch, pw + w))
        elif toks and toks[-1][0] and toks[-1][0][-1] in _NO_BREAK_AFTER:
            prev, pw = toks.pop()
            toks.append((prev + ch, pw + w))
        else:
            toks.append((ch, w))
    flush()
    if mode == "lo":  # 中西文（含空格、标点、运算符）交界自动间距，约 0.25em
        out = []
        for i, (t, w) in enumerate(toks):
            if i > 0 and out:
                a, b = out[-1][0][-1], t[0]
                if (ord(a) > 127) != (ord(b) > 127):
                    w += AUTOSPACE
            out.append((t, w))
        toks = out
    return toks


def wrap_stats(text, size_pt, width_pt, bold=False, mode="ppt"):
    """贪心折行；返回 (行数, 末行宽度pt)。"""
    if not text:
        return 1, 0.0
    lines, x = 1, 0.0
    for tok, w in _tokens(text, bold, mode):
        wp = w * size_pt
        if tok == " ":
            if x == 0:
                continue
            x += wp
            continue
        if x + wp <= width_pt + 0.5:
            x += wp
            continue
        if wp > width_pt:  # 超长西文串：逐字符断开
            for ch in tok:
                cwp = char_em(ch, mode) * size_pt
                if x + cwp > width_pt + 0.5:
                    lines += 1
                    x = 0.0
                x += cwp
            continue
        lines += 1
        x = wp
    return lines, x


# --------------------------------------------------------------------------
# 行内标记
# --------------------------------------------------------------------------
_INL = re.compile(r"\[\[(.+?)\]\]|\*\*(.+?)\*\*|\^\{(.+?)\}|_\{(.+?)\}|`(.+?)`", re.S)


def parse_inline(s, bold=False, color=None, base=0, mono=False):
    """返回 [(text, bold, color, baseline, mono)]"""
    out = []
    pos = 0
    for m in _INL.finditer(s):
        if m.start() > pos:
            out.append((s[pos:m.start()], bold, color, base, mono))
        if m.group(1) is not None:
            out += parse_inline(m.group(1), True, RED, base, mono)
        elif m.group(2) is not None:
            out += parse_inline(m.group(2), True, color, base, mono)
        elif m.group(3) is not None:
            out += parse_inline(m.group(3), bold, color, 30000, mono)
        elif m.group(4) is not None:
            out += parse_inline(m.group(4), bold, color, -25000, mono)
        else:
            out += parse_inline(m.group(5), bold, color, base, True)
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], bold, color, base, mono))
    return out


def plain_text(s):
    """去掉标记后的纯文本（用于度量）。"""
    return "".join(t for t, *_ in parse_inline(s))


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def run_xml(text, sz, bold=False, color=None, base=0, mono=False, latin_only_font=None):
    attrs = f'lang="zh-CN" altLang="en-US" sz="{int(round(sz * 100))}"'
    if bold:
        attrs += ' b="1"'
    if base:
        attrs += f' baseline="{base}"'
    attrs += ' dirty="0"'
    fill = f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>' if color else ""
    latin = MONO if mono else LATIN
    return (f'<a:r><a:rPr {attrs}>{fill}<a:latin typeface="{latin}"/>'
            f'<a:ea typeface="{EA}"/></a:rPr><a:t>{esc(text)}</a:t></a:r>')


def end_rpr(sz):
    return (f'<a:endParaRPr lang="zh-CN" altLang="en-US" sz="{int(round(sz * 100))}" dirty="0">'
            f'<a:latin typeface="{LATIN}"/><a:ea typeface="{EA}"/></a:endParaRPr>')


# --------------------------------------------------------------------------
# 段落
# --------------------------------------------------------------------------
class Para:
    """一个段落：kind 决定项目符号/粗体/颜色。"""

    KINDS = ("p", "h", "b1", "hb", "b2", "red", "sp")

    def __init__(self, kind, text):
        self.kind = kind
        self.text = text

    @staticmethod
    def parse(it):
        if isinstance(it, Para):
            return it
        if isinstance(it, tuple):
            return Para(it[0], it[1])
        if it == "":
            return Para("sp", "")
        for pre, kind in (("#- ", "hb"), ("-- ", "b2"), ("- ", "b1"),
                          ("# ", "h"), ("!! ", "red")):
            if it.startswith(pre):
                return Para(kind, it[len(pre):])
        return Para("p", it)

    # 缩进（EMU）
    def marL(self):
        return {"b1": 342900, "hb": 342900, "b2": 742950}.get(self.kind, 0)

    def size(self, S):
        return SUB_SIZE.get(S, max(14, S - 4)) if self.kind == "b2" else S

    def xml(self, S, spc_aft):
        sz = self.size(S)
        k = self.kind
        if k == "sp":
            return (f'<a:p><a:pPr marL="0" indent="0" algn="l"><a:spcBef><a:spcPts val="0"/></a:spcBef>'
                    f'<a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>{end_rpr(max(8, S // 2))}</a:p>')
        if k in ("b1", "hb"):
            ppr = ('<a:pPr marL="342900" indent="-342900" algn="l" eaLnBrk="1" fontAlgn="auto" hangingPunct="1">'
                   f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{int(spc_aft * 100)}"/></a:spcAft>'
                   '<a:buClr><a:srgbClr val="000000"/></a:buClr><a:buFont typeface="Arial"/><a:buChar char="•"/></a:pPr>')
        elif k == "b2":
            ppr = ('<a:pPr marL="742950" lvl="1" indent="-285750" algn="l" eaLnBrk="1" fontAlgn="auto" hangingPunct="1">'
                   f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{int(spc_aft * 100)}"/></a:spcAft>'
                   '<a:buClr><a:srgbClr val="000000"/></a:buClr><a:buFont typeface="Arial"/><a:buChar char="–"/></a:pPr>')
        else:
            ppr = ('<a:pPr marL="0" indent="0" algn="l" eaLnBrk="1" fontAlgn="auto" hangingPunct="1">'
                   f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{int(spc_aft * 100)}"/></a:spcAft>'
                   '<a:buNone/></a:pPr>')
        base_bold = k in ("h", "hb", "red")
        base_color = RED if k == "red" else None
        runs = "".join(run_xml(t, sz, b, c, bl, mo)
                       for t, b, c, bl, mo in parse_inline(self.text, base_bold, base_color))
        return f"<a:p>{ppr}{runs}{end_rpr(sz)}</a:p>"

    def height_pt(self, S, width_pt, spc_aft, mode="ppt"):
        """估算本段高度（pt）、末行宽度、行数、字号。"""
        sz = self.size(S)
        if self.kind == "sp":
            return max(8, S // 2) * LINE_K, 0.0, 1, sz
        w = width_pt - self.marL() / EMU_PT
        txt = plain_text(self.text)
        bold = self.kind in ("h", "hb", "red") or "[[" in self.text or "**" in self.text
        global _MONO_ALL
        _MONO_ALL = "`" in self.text
        try:
            n, last = wrap_stats(txt, sz, w, bold, mode)
        finally:
            _MONO_ALL = False
        return n * sz * LINE_K + spc_aft, last, n, sz


def measure(paras, S, width_emu, spc_aft):
    """文本框内容总高度（pt，两种度量模式取大）与孤字警告列表。"""
    width_pt = (width_emu - 2 * INS_LR) / EMU_PT
    worst = 0.0
    orphans = []
    for mode in MODES:
        total = 0.0
        for p in paras:
            h, last, n, sz = p.height_pt(S, width_pt, spc_aft, mode)
            total += h
            if n > 1 and 0 < last < 2.2 * sz and p.kind != "sp" and p.text[:18] not in orphans:
                orphans.append(p.text[:18])
        worst = max(worst, total)
    return worst, orphans


def fit_size(paras, width_emu, height_emu, spc_aft=6, sizes=(28, 26, 24, 22, 20, 18, 16), slack=0.04,
             orphan_steps=2):
    """选最大且放得下的字号；若相邻较小字号（最多 orphan_steps 档）可避免孤字行，则取之。"""
    avail = (height_emu - 2 * INS_TB) / EMU_PT * (1 - slack)
    fits = []
    last = None
    for S in sizes:
        h, orphans = measure(paras, S, width_emu, spc_aft)
        last = (S, h, avail, orphans)
        if h <= avail:
            fits.append(last)
    if not fits:
        return last
    first = fits[0]
    for cand in fits[:orphan_steps + 1]:
        if not cand[3]:
            return cand
    return first


# --------------------------------------------------------------------------
# 形状 XML
# --------------------------------------------------------------------------
def title_xml(shape_id, text, sz):
    return (f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{shape_id}" name="Title {shape_id - 1}"/>'
            '<p:cNvSpPr><a:spLocks noGrp="1" noRot="1" noChangeArrowheads="1"/></p:cNvSpPr>'
            '<p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="15033" y="304882"/><a:ext cx="9108305" cy="457200"/></a:xfrm></p:spPr>'
            '<p:txBody><a:bodyPr rtlCol="0"><a:normAutofit fontScale="90000"/></a:bodyPr><a:lstStyle/>'
            '<a:p><a:pPr algn="l" eaLnBrk="1" fontAlgn="auto" hangingPunct="1"><a:spcAft><a:spcPts val="0"/></a:spcAft></a:pPr>'
            f'<a:r><a:rPr lang="zh-CN" altLang="en-US" sz="{sz}" b="1" dirty="0">'
            f'<a:latin typeface="{EA}"/><a:ea typeface="{EA}"/></a:rPr><a:t>{esc(text)}</a:t></a:r></a:p>'
            '</p:txBody></p:sp>')


def textbox_xml(shape_id, name, x, y, cx, cy, paras_xml, anchor="ctr"):
    return (f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/>'
            '<p:cNvSpPr txBox="1"><a:spLocks noRot="1" noChangeArrowheads="1"/></p:cNvSpPr><p:nvPr/></p:nvSpPr>'
            f'<p:spPr bwMode="auto"><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(cx)}" cy="{int(cy)}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
            '<p:txBody><a:bodyPr vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" '
            f'numCol="1" rtlCol="0" anchor="{anchor}" anchorCtr="0" compatLnSpc="1">'
            '<a:prstTxWarp prst="textNoShape"><a:avLst/></a:prstTxWarp><a:noAutofit/></a:bodyPr><a:lstStyle/>'
            f'{paras_xml}</p:txBody></p:sp>')


def _ln(tag):
    return (f'<a:{tag} w="6350"><a:solidFill><a:srgbClr val="{GRID}"/></a:solidFill></a:{tag}>')


def cell_xml(text, sz, header=False, bold=False, fill="FFFFFF", align="l"):
    paras = []
    for line in str(text).split("\n"):
        runs = []
        for t, b, c, bl, mo in parse_inline(line, bold or header, "FFFFFF" if header else None):
            runs.append(run_xml(t, sz, b, c, bl, mo))
        paras.append(f'<a:p><a:pPr algn="{align}"><a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
                     f'{"".join(runs)}<a:endParaRPr lang="zh-CN" altLang="en-US" sz="{int(sz * 100)}"/></a:p>')
    return ('<a:tc><a:txBody><a:bodyPr/><a:lstStyle/>' + "".join(paras) + '</a:txBody>'
            '<a:tcPr marL="72000" marR="72000" anchor="ctr">'
            + _ln("lnL") + _ln("lnR") + _ln("lnT") + _ln("lnB")
            + f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill></a:tcPr></a:tc>')


def table_xml(shape_id, x, y, widths, heights, header, rows, sz, first_bold=True, aligns=None):
    ncol = len(widths)
    aligns = aligns or ["l"] * ncol
    grid = "".join(f'<a:gridCol w="{int(w)}"/>' for w in widths)
    trs = []
    trs.append(f'<a:tr h="{int(heights[0])}">' + "".join(
        cell_xml(h, sz, header=True, fill=NAVY, align=aligns[i]) for i, h in enumerate(header)) + "</a:tr>")
    for ri, row in enumerate(rows):
        fill = "FFFFFF" if ri % 2 == 0 else BAND
        trs.append(f'<a:tr h="{int(heights[ri + 1])}">' + "".join(
            cell_xml(c, sz, bold=(first_bold and ci == 0), fill=fill, align=aligns[ci])
            for ci, c in enumerate(row)) + "</a:tr>")
    total_w = int(sum(widths))
    total_h = int(sum(heights))
    return (f'<p:graphicFrame {NSDECL}><p:nvGraphicFramePr><p:cNvPr id="{shape_id}" name="Table {shape_id}"/>'
            '<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
            f'<p:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{total_w}" cy="{total_h}"/></p:xfrm>'
            '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
            f'<a:tbl><a:tblPr firstRow="1" bandRow="1"/><a:tblGrid>{grid}</a:tblGrid>{"".join(trs)}</a:tbl>'
            '</a:graphicData></a:graphic></p:graphicFrame>')


# --------------------------------------------------------------------------
# 标题字号
# --------------------------------------------------------------------------
def title_size(text):
    """40pt（×90% = 36pt 显示）能放进一行就用 4000，否则逐档缩小。"""
    width_pt = (9108305 - 2 * INS_LR) / EMU_PT
    w_em = sum(char_em(ch) * (1.04 if ord(ch) < 0x80 else 1.0) for ch in text)
    for sz in (4000, 3800, 3600, 3400, 3200, 3000, 2800, 2600):
        if w_em * sz / 100 * 0.9 <= width_pt:
            return sz
    return 2600


# --------------------------------------------------------------------------
# Deck
# --------------------------------------------------------------------------
class Deck:
    def __init__(self, template=TEMPLATE, assets=ASSETS):
        self.prs = Presentation(template)
        self.assets = assets
        self.warnings = []
        self.count = 1  # 封面算第 1 页
        self.log = []   # [(页号, 标题, 类型, 字号)]
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst)[1:]:
            sldIdLst.remove(sldId)
            self.prs.part.drop_rel(sldId.rId)
        self.layout = next(l for l in self.prs.slide_layouts if l.name == "标题和内容")

    # ---------------- 公共 -----------------
    def warn(self, msg):
        self.warnings.append(f"[p{self.count}] {msg}")

    def _new(self, title):
        slide = self.prs.slides.add_slide(self.layout)
        tree = slide.shapes._spTree
        for sp in list(tree.findall(qn("p:sp"))):
            tree.remove(sp)
        self.count += 1
        sz = title_size(title)
        tree.append(etree.fromstring(title_xml(2, title, sz)))
        self._tree = tree
        return slide, tree

    def _id(self):
        """下一个未被占用的形状 id（图片由 python-pptx 另行分配，需避开）。"""
        ids = [int(e.get("id")) for e in self._tree.iter()
               if e.tag.endswith("}cNvPr") and str(e.get("id", "")).isdigit()]
        return max(ids + [1]) + 1

    def _notes(self, slide, notes):
        if not notes:
            return
        if isinstance(notes, (list, tuple)):
            notes = "\n".join(notes)
        slide.notes_slide.notes_text_frame.text = normalize_notes(notes)

    def _img(self, name):
        p = name if os.path.isabs(name) else os.path.join(self.assets, name)
        if not os.path.exists(p):
            raise FileNotFoundError(p)
        with Image.open(p) as im:
            return p, im.size

    # ---------------- 封面 -----------------
    def cover(self, title, notes=None):
        slide = self.prs.slides[0]
        for t in slide.shapes._spTree.iter(qn("a:t")):
            if t.text and "章" in t.text:
                t.text = title
                break
        if notes:
            self._notes(slide, notes)
        self.log.append((1, title, "cover", 0, "1"))

    # ---------------- 文字页 -----------------
    def text(self, title, items, notes=None, size=None, max_size=26, spc=6, anchor="ctr",
             x=BODY_X, y=BODY_Y, w=BODY_W, h=BODY_H, src=None):
        slide, tree = self._new(title)
        paras = [Para.parse(i) for i in items]
        sizes = tuple(s for s in (28, 26, 24, 22, 20, 18, 16) if s <= max_size)
        if size:
            S = size
            tot, orphans = measure(paras, S, w, spc)
            avail = (h - 2 * INS_TB) / EMU_PT
        else:
            S, tot, avail, orphans = fit_size(paras, w, h, spc, sizes)
        if tot > avail:
            self.warn(f"文字溢出：{title}（估计{tot:.0f}pt > 可用{avail:.0f}pt，字号{S}）")
        if orphans:
            self.warn(f"孤字行：{title} → {orphans}")
        pxml = "".join(p.xml(S, spc) for p in paras)
        tree.append(etree.fromstring(textbox_xml(self._id(), "Rectangle 2", x, y, w, h, pxml, anchor)))
        self._notes(slide, notes)
        self.log.append((self.count, title, "text", S, src))
        return slide

    # ---------------- 表格页 -----------------
    def table(self, title, header, rows, widths=None, foot=None, notes=None, fs=None,
              first_bold=True, aligns=None, x=152516, y=950000, total_w=8800000,
              foot_size=20, min_row=450000, lead=None, lead_size=20, src=None):
        """表格页。lead：表格上方的导语（段落列表）；foot：表格下方的红色粗体收束句（字符串或列表）。"""
        slide, tree = self._new(title)
        ncol = len(header)
        rows = [list(r) for r in rows]
        for r in rows:
            assert len(r) == ncol, f"表格列数不一致：{title} {r}"
        # 列宽
        if widths is None:
            nat = []
            for ci in range(ncol):
                m = max(sum(char_em(ch) for ch in plain_text(line))
                        for r in [header] + rows for line in str(r[ci]).split("\n"))
                nat.append(max(m, 3.0))
            weights = [n ** 0.72 for n in nat]
        else:
            weights = list(widths)
        sw = float(sum(weights))
        cw = [total_w * v / sw for v in weights]
        box_w = total_w + 39200
        inner_pt = (box_w - 2 * INS_LR) / EMU_PT * 0.97  # lead/foot 略保守，避免估算偏小压住表格
        lead_paras = [Para.parse(i) for i in (lead or [])]
        if isinstance(foot, str):
            foot = [foot]
        foot_paras = []
        for f in (foot or []):
            pp = Para.parse(f)
            if pp.kind == "p" and not str(f).startswith(("#", "-", "!")):
                pp = Para("red", f)
            foot_paras.append(pp)
        lead_h = 0
        if lead_paras:
            lead_h = max(sum(q.height_pt(lead_size, inner_pt, 4, m)[0] for q in lead_paras) for m in MODES) * EMU_PT \
                + 2 * INS_TB + 60000
        sizes = [fs] if fs else [18, 17, 16, 15, 14, 13, 12]
        chosen = None
        cw0 = list(cw)
        for sz in sizes:
            cw = self._opt_widths(cw0, header, rows, sz, first_bold, total_w)
            heights = self._row_heights(header, rows, cw, sz, min_row, first_bold)
            tbl_h = sum(heights)
            foot_h = 0
            if foot_paras:
                foot_h = max(sum(q.height_pt(foot_size, inner_pt, 4, m)[0] for q in foot_paras) for m in MODES) \
                    * EMU_PT + 2 * INS_TB + 200000
            bottom = y + lead_h + tbl_h + foot_h
            chosen = (sz, heights, tbl_h, foot_h, bottom, cw)
            if bottom <= 6700000:
                break
        sz, heights, tbl_h, foot_h, bottom, cw = chosen
        if bottom > 6700000:
            self.warn(f"表格溢出：{title}（底边{bottom / 12700:.0f}pt，字号{sz}）")
        if lead_paras:
            pxml = "".join(q.xml(lead_size, 4) for q in lead_paras)
            tree.append(etree.fromstring(textbox_xml(self._id(), "TextBox 5", x, y, box_w, lead_h - 60000, pxml, "t")))
        tree.append(etree.fromstring(table_xml(self._id(), x, y + lead_h, cw, heights, header, rows, sz,
                                               first_bold, aligns)))
        if foot_paras:
            pxml = "".join(q.xml(foot_size, 4) for q in foot_paras)
            tree.append(etree.fromstring(textbox_xml(self._id(), "TextBox 7", x, y + lead_h + tbl_h + 200000,
                                                     box_w, foot_h - 200000, pxml, "t")))
        self._notes(slide, notes)
        self.log.append((self.count, title, "table", sz, src))
        return slide

    def _table_cost(self, cw, header, rows, sz, first_bold):
        """列宽方案的代价：单元格孤字行×1000 + 各行高度之和(pt)。"""
        orph = 0
        total = 0.0
        for ri, row in enumerate([header] + rows):
            hmax = 0.0
            for ci, cell in enumerate(row):
                width_pt = (cw[ci] - 2 * 72000) / EMU_PT
                bold = ri == 0 or (ci == 0 and first_bold)
                lines = 0
                bad = False
                for line in str(cell).split("\n"):
                    wn = 0
                    for m in MODES:
                        n, last = wrap_stats(plain_text(line), sz, width_pt, bold, m)
                        wn = max(wn, n)
                        if n > 1 and 0 < last < 2.2 * sz:
                            bad = True
                    lines += wn
                orph += 1 if bad else 0
                hmax = max(hmax, lines * sz * LINE_K + 8.2)
            total += hmax
        return orph * 1000 + total, orph

    def _opt_widths(self, cw, header, rows, sz, first_bold, total_w):
        """在原列宽附近微调（每列 ≥ 最长不可断词，且不偏离原宽 ±40%），尽量消除单元格孤字行并压低表高。"""
        n = len(cw)
        base = list(cw)
        lo = []
        for ci in range(n):
            tokw = 3.0
            for ri, row in enumerate([header] + rows):
                for line in str(row[ci]).split("\n"):
                    for t, w in _tokens(plain_text(line), True, "lo"):
                        tokw = max(tokw, w)
            lo.append(max(tokw * sz * EMU_PT + 2 * 72000 + 30000, base[ci] * 0.6))
        hi = [b * 1.4 for b in base]
        cur = list(base)
        best, borph = self._table_cost(cur, header, rows, sz, first_bold)
        if borph == 0:
            return cur
        step = total_w * 0.015
        for _ in range(60):
            improved = False
            cand_best = None
            for i in range(n):
                for j in range(n):
                    if i == j:
                        continue
                    for k in (1, 2, 4, 6):
                        d = step * k
                        if cur[j] - d < lo[j] or cur[i] + d > hi[i]:
                            continue
                        trial = list(cur)
                        trial[i] += d
                        trial[j] -= d
                        c, o = self._table_cost(trial, header, rows, sz, first_bold)
                        if c < best - 0.01 and (cand_best is None or c < cand_best[0]):
                            cand_best = (c, o, trial)
            if cand_best:
                best, borph, cur = cand_best
                improved = True
            if not improved or borph == 0:
                break
        return cur

    def _row_heights(self, header, rows, cw, sz, min_row, first_bold=True):
        out = []
        for ri, row in enumerate([header] + rows):
            hmax = 0.0
            for ci, cell in enumerate(row):
                width_pt = (cw[ci] - 2 * 72000) / EMU_PT
                lines = 0
                for line in str(cell).split("\n"):
                    n = max(wrap_stats(plain_text(line), sz, width_pt, ri == 0 or (ci == 0 and first_bold), m)[0]
                            for m in MODES)
                    lines += n
                hh = lines * sz * LINE_K + 7.2 + 1.0
                hmax = max(hmax, hh)
            out.append(max(min_row, hmax * EMU_PT))
        return out

    # ---------------- 图片页 -----------------
    def pic(self, title, img, items=None, notes=None, layout="top", max_size=24, spc=6,
            alt=None, text_anchor=None, src=None):
        imgs = [img] if isinstance(img, str) else list(img)
        meta = [self._img(n) for n in imgs]
        paras = [Para.parse(i) for i in (items or [])]
        slide, tree = self._new(title)
        ar = sum(sz[0] / sz[1] for _, sz in meta)  # 同高并排时 总宽/高
        gap_x = 150000
        Y0, Y1 = 950000, 6700000
        if layout == "top":
            WMAX = 8500000 - (len(meta) - 1) * gap_x
            h_by_w = WMAX / ar
            chosen = None
            if not paras:
                H = min(h_by_w, Y1 - Y0)
                chosen = (0, H, 0)
            else:
                cands, last = [], None
                for S in [s for s in (24, 22, 20, 18, 16) if s <= max_size]:
                    tot, orph = measure(paras, S, 8839200, spc)
                    cap_h = tot * EMU_PT + 2 * INS_TB + 150000
                    H = min(h_by_w, (Y1 - Y0) - cap_h)
                    last = (S, H, cap_h, orph)
                    if H >= min(h_by_w, 2300000):
                        cands.append(last)
                if cands:
                    pick = next((c for c in cands[:3] if not c[3]), cands[0])
                else:
                    pick = last
                chosen = pick[:3]
                if chosen[1] < 1500000:
                    self.warn(f"图片偏小：{title}（高{chosen[1] / 12700:.0f}pt），考虑缩短文字或改用 left/right")
            S, H, cap_h = chosen
            Wtot = H * ar + (len(meta) - 1) * gap_x
            xx = (SLIDE_W - Wtot) / 2
            for (p, (iw, ih)), nm in zip(meta, imgs):
                w = H * iw / ih
                pic = slide.shapes.add_picture(p, Emu(int(xx)), Emu(int(Y0)), Emu(int(w)), Emu(int(H)))
                pic._element.nvPicPr.cNvPr.set("descr", alt or os.path.basename(nm))
                xx += w + gap_x
            if paras:
                ty = Y0 + H + 150000
                th = Y1 - ty + 0
                tot, orphans = measure(paras, S, 8839200, spc)
                if tot > (th - 2 * INS_TB) / EMU_PT:
                    self.warn(f"图片页文字溢出：{title}")
                if orphans:
                    self.warn(f"孤字行：{title} → {orphans}")
                pxml = "".join(p.xml(S, spc) for p in paras)
                tree.append(etree.fromstring(textbox_xml(self._id(), "Rectangle 2", 152516, ty, 8839200, th,
                                                         pxml, text_anchor or "t")))
        elif layout in ("left", "right"):
            iw, ih = meta[0][1]
            IMG_WMAX = 3900000 if ih >= iw else 4300000
            IMG_HMAX = Y1 - Y0
            scale = min(IMG_WMAX / iw, IMG_HMAX / ih)
            w, h = iw * scale, ih * scale
            tw = 8839200 - w - 200000
            S, tot, avail, orphans = fit_size(paras, tw, Y1 - Y0, spc,
                                              [s for s in (24, 22, 20, 18, 16) if s <= max_size])
            if tot > avail:
                self.warn(f"图片页文字溢出：{title}（估计{tot:.0f}>{avail:.0f}，字号{S}）")
            if orphans:
                self.warn(f"孤字行：{title} → {orphans}")
            if layout == "left":
                ix, tx = 152516, 152516 + w + 200000
            else:
                tx, ix = 152516, 152516 + tw + 200000
            iy = Y0 + (Y1 - Y0 - h) / 2
            pic = slide.shapes.add_picture(meta[0][0], Emu(int(ix)), Emu(int(iy)), Emu(int(w)), Emu(int(h)))
            pic._element.nvPicPr.cNvPr.set("descr", alt or os.path.basename(imgs[0]))
            pxml = "".join(p.xml(S, spc) for p in paras)
            tree.append(etree.fromstring(textbox_xml(self._id(), "Rectangle 2", tx, Y0, tw, Y1 - Y0,
                                                     pxml, text_anchor or "ctr")))
        else:
            raise ValueError(layout)
        self._notes(slide, notes)
        self.log.append((self.count, title, "pic:" + layout, S, src))
        return slide

    # ---------------- 保存 -----------------
    def save(self, path, title="数字货币 第四章 比特币"):
        cp = self.prs.core_properties
        cp.title = title
        cp.modified = datetime.datetime.now()
        self.prs.save(path)
        return path


def normalize_notes(text):
    """备注风格：小标题独占一行、正文另起一段（与 3.pptx 一致）。
    源文本可写成“【讲解】正文”，转换为“讲解\n正文”。"""
    t = text.strip()
    t = re.sub(r"(?m)^[ \t]*【([^】\n]{1,16})】[ \t]*", lambda m: "\n" + m.group(1) + "\n", t)
    t = re.sub(r"\n{2,}", "\n", t).strip()
    return t


def report(d, out=sys.stdout):
    for w in d.warnings:
        print("WARN", w, file=out)
    print(f"slides={d.count}  warnings={len(d.warnings)}", file=out)
