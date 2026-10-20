#!/usr/bin/env python3
"""finalize.py deck.pptx [thumb_source.pdf] — 修正模板遗留的元数据：
  · docProps/app.xml：幻灯片数/备注数/字数，并去掉过期的“幻灯片标题”列表；
  · docProps/thumbnail.jpeg：换成本课件封面（否则资源管理器里仍显示“第三章”）。
幂等；就地改写。"""
import os, re, shutil, subprocess, sys, tempfile, zipfile
from pptx import Presentation

path = os.path.abspath(sys.argv[1])
pdf = os.path.abspath(sys.argv[2]) if len(sys.argv) > 2 else None
SOFFICE = "/mnt/skills/public/pptx/scripts/office/soffice.py"

prs = Presentation(path)
n = len(prs.slides)
notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
paras = words = 0
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.text.strip():
                    paras += 1
                    words += len(p.text.replace(" ", ""))

tmp = tempfile.mkdtemp(prefix="fin_")
thumb = None
try:
    if pdf is None:
        subprocess.run(["python3", SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    if os.path.exists(pdf):
        subprocess.run(["pdftoppm", "-jpeg", "-f", "1", "-l", "1", "-scale-to-x", "256", "-scale-to-y", "192",
                        pdf, os.path.join(tmp, "thumb")], check=True)
        cands = [f for f in os.listdir(tmp) if f.startswith("thumb") and f.endswith(".jpg")]
        if cands:
            thumb = open(os.path.join(tmp, cands[0]), "rb").read()

    out = os.path.join(tmp, "new.pptx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                a = data.decode("utf8")
                a = re.sub(r"<Slides>\d+</Slides>", f"<Slides>{n}</Slides>", a)
                a = re.sub(r"<Notes>\d+</Notes>", f"<Notes>{notes}</Notes>", a)
                a = re.sub(r"<Words>\d+</Words>", f"<Words>{words}</Words>", a)
                a = re.sub(r"<Paragraphs>\d+</Paragraphs>", f"<Paragraphs>{paras}</Paragraphs>", a)
                a = re.sub(r"<TotalTime>\d+</TotalTime>", "<TotalTime>0</TotalTime>", a)
                a = re.sub(r"<HeadingPairs>.*?</HeadingPairs>", "", a, flags=re.S)
                a = re.sub(r"<TitlesOfParts>.*?</TitlesOfParts>", "", a, flags=re.S)
                data = a.encode("utf8")
            elif item.filename == "docProps/thumbnail.jpeg" and thumb:
                data = thumb
            zout.writestr(item, data)
    shutil.move(out, path)
finally:
    shutil.rmtree(tmp, ignore_errors=True)
print(f"finalize: slides={n} notes={notes} paras={paras} chars={words} thumb={'yes' if thumb else 'no'}")
