#!/usr/bin/env python3
"""lint.py deck.pptx —— 内容层面的自动检查：残留标记、页码引用、emoji、页上“来源”小字、过长标题、空备注、术语一致性。"""
import re, sys, collections
from pptx import Presentation

prs = Presentation(sys.argv[1])
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⭐⭕✅❌]")
issues = collections.defaultdict(list)

def add(i, msg):
    issues[i].append(msg)

titles = collections.Counter()
for i, s in enumerate(prs.slides, 1):
    title = s.shapes.title.text_frame.text if s.shapes.title else ""
    titles[title] += 1
    if i > 1 and len(title) > 22:
        add(i, f"标题偏长({len(title)}字)：{title}")
    body = []
    for sh in s.shapes:
        if sh.has_text_frame:
            body.append(sh.text_frame.text)
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    body.append(c.text_frame.text)
    text = "\n".join(body)
    notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    for label, t in (("页面", text), ("备注", notes)):
        if re.search(r"\[\[|\]\]|\*\*|\^\{|_\{|(?m:^!! )|(?m:^-- )|(?m:^#- )", t):
            add(i, f"{label}含未解析的标记")
        if EMOJI.search(t):
            add(i, f"{label}含 emoji：{EMOJI.search(t).group(0)}")
        m = re.search(r"第\s*\d+\s*页", t)
        if m and "白皮书" not in t[max(0, m.start() - 12):m.end() + 4] and "原课件" not in t[max(0, m.start() - 12):m.end() + 4] and not re.search(r"原[PＰ]|P\d", t[max(0, m.start() - 8):m.end()]):
            add(i, f"{label}有页码引用：{m.group(0)}")
        if re.search(r"(?m:^来源[：:])", t) and label == "页面":
            add(i, "页面上有“来源：”小字（应放入备注）")
        if "匿名" in t and "假名" not in t and "不是匿名" not in t and "并非匿名" not in t and "匿名的" not in t and i > 1:
            add(i, f"{label}提到“匿名”，请确认应为“假名”")
    if i > 1 and not notes.strip():
        add(i, "备注为空")
    if i > 1 and len(text.replace("\n", "").strip()) < 8 and not any(sh.shape_type == 13 for sh in s.shapes):
        add(i, "页面几乎没有文字")
dups = [t for t, n in titles.items() if n > 1 and t]
n_issue = sum(len(v) for v in issues.values())
for i in sorted(issues):
    for m in issues[i]:
        print(f"p{i}: {m}")
if dups:
    print("重复标题：", "；".join(f"{t}×{titles[t]}" for t in dups))
print(f"slides={len(prs.slides)} issues={n_issue}")
